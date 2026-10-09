from __future__ import annotations

import argparse
import csv
import json
import re
import statistics
import time
from collections import defaultdict
from datetime import datetime, timezone, timedelta, date
from zoneinfo import ZoneInfo
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
try:
    from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
    from playwright.sync_api import sync_playwright
except ImportError:  # Lets --help / analytics code fail with a useful setup message.
    PlaywrightTimeoutError = TimeoutError
    sync_playwright = None


ROOT = Path(__file__).resolve().parent
RAW_FIELDS = [
    "match_id", "scraped_at_utc", "league_url", "country", "league", "season",
    "date_time", "game_date", "status", "home_team", "away_team", "home_total", "away_total",
    "home_q1", "away_q1", "home_q2", "away_q2", "home_q3", "away_q3",
    "home_q4", "away_q4", "home_ot", "away_ot", "notes", "tags",
]


def text_or_blank(row, selector: str) -> str:
    loc = row.locator(selector).first
    if loc.count() == 0:
        return ""
    try:
        return loc.inner_text(timeout=1000).strip()
    except PlaywrightTimeoutError:
        return ""


def number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    m = re.search(r"-?\d+(?:\.\d+)?", str(value))
    return float(m.group()) if m else None


def clean_match_id(raw: str) -> str:
    # Flashscore event row ids normally look like g_3_MATCHID.
    return raw.rsplit("_", 1)[-1] if raw else ""


def page_context(page, url: str) -> tuple[str, str, str]:
    country = text_or_blank(page, ".breadcrumb__link:nth-last-child(2)")
    league = text_or_blank(page, ".heading__name") or text_or_blank(page, ".tournamentHeader__country")
    season = text_or_blank(page, ".heading__info")
    if not league:
        parts = [p for p in url.split("/") if p]
        league = parts[-2] if len(parts) >= 2 else "unknown"
    return country, league, season


def parse_row(row, url: str, context: tuple[str, str, str]) -> dict[str, Any] | None:
    country, league, season = context
    home = text_or_blank(row, ".event__homeParticipant, .event__participant--home")
    away = text_or_blank(row, ".event__awayParticipant, .event__participant--away")
    if not home or not away:
        raise RuntimeError("Could not read game teams; Flashscore markup changed. CSV left unchanged.")

    values: dict[str, Any] = {
        "match_id": clean_match_id(row.get_attribute("id") or ""),
        "scraped_at_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "league_url": url,
        "country": country,
        "league": league,
        "season": season,
        "date_time": text_or_blank(row, ".event__stageTime--date, .event__time"),
        "status": text_or_blank(row, ".event__stage--block") or "Finished",
        "home_team": home,
        "away_team": away,
        "home_total": text_or_blank(row, ".event__score--home"),
        "away_total": text_or_blank(row, ".event__score--away"),
        "notes": "",
        "tags": "",
    }
    for q in range(1, 5):
        values[f"home_q{q}"] = text_or_blank(row, f".event__part--home.event__part--{q}")
        values[f"away_q{q}"] = text_or_blank(row, f".event__part--away.event__part--{q}")

    # OT, when displayed, is the fifth partial score. Keep it separate from Q4.
    values["home_ot"] = text_or_blank(row, ".event__part--home.event__part--5")
    values["away_ot"] = text_or_blank(row, ".event__part--away.event__part--5")

    if not values["match_id"]:
        values["match_id"] = "|".join([league, values["date_time"], home, away])
    return values


def result_date(value: str, today: date) -> date | None:
    match = re.search(r"(\d{1,2})\.(\d{1,2})\.(?:(\d{4})(?!\d))?", value)
    if not match:
        return None
    day, month, year = match.groups()
    try:
        parsed = date(int(year) if year else today.year, int(month), int(day))
        if not year and parsed > today:
            parsed = parsed.replace(year=today.year - 1)
        return parsed
    except ValueError:
        return None


def scrape(config: dict[str, Any]) -> list[dict[str, Any]]:
    if sync_playwright is None:
        raise SystemExit("Playwright is not installed. Run: py -m pip install -r requirements.txt")
    today = datetime.now(ZoneInfo(config.get("timezone", "Europe/Budapest"))).date()
    start = today - timedelta(days=int(config.get("lookback_days", 3)))
    print(f"Date window: {start} through {today - timedelta(days=1)}")
    results: list[dict[str, Any]] = []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=bool(config.get("headless", True)))
        page = browser.new_page(viewport={"width": 1440, "height": 1100})
        for url in config["league_urls"]:
            print(f"Scraping {url}")
            page.goto(url, wait_until="domcontentloaded", timeout=60000)
            try:
                page.locator(".event__match").first.wait_for(timeout=20000)
            except PlaywrightTimeoutError:
                raise RuntimeError(f"No match rows found for {url}; CSV left unchanged.")

            consent = page.locator("#onetrust-reject-all-handler")
            if consent.count() and consent.is_visible():
                consent.click()

            for _ in range(int(config.get("max_show_more_clicks", 3))):
                more = page.get_by_text(re.compile(r"show more matches", re.I)).first
                if more.count() == 0 or not more.is_visible():
                    break
                more.click()
                page.wait_for_timeout(1000)

            context = page_context(page, url)
            rows = page.locator(".event__match")
            print(f"  Found {rows.count()} game rows")
            collected_before = len(results)
            for i in range(rows.count()):
                item = parse_row(rows.nth(i), url, context)
                if item:
                    played = result_date(item["date_time"], today)
                    if played is None:
                        raise RuntimeError(f"Unrecognized game date: {item['date_time']!r}; CSV left unchanged.")
                    if start <= played < today:
                        scores = [item[f"{side}_q{q}"] for side in ("home", "away") for q in range(1, 5)]
                        if not all(re.fullmatch(r"\d+", str(v)) for v in scores):
                            raise RuntimeError(f"Missing quarter scores for {item['match_id']}; CSV left unchanged.")
                        for side in ("home", "away"):
                            total = number(item[f"{side}_total"])
                            regulation = sum(int(item[f"{side}_q{q}"]) for q in range(1,5))
                            if total is None or total < regulation:
                                raise RuntimeError("Invalid final score; CSV left unchanged.")
                            item[f"{side}_ot"] = int(total - regulation)
                        item["game_date"] = played.isoformat()
                        results.append(item)
            print(f"  Collected {len(results) - collected_before} games in date window")
            time.sleep(float(config.get("request_delay_seconds", 1.5)))
        browser.close()
    return results


def read_existing_csv(path: Path) -> dict[str, dict[str, str]]:
    if not path.exists():
        return {}
    with path.open("r", newline="", encoding="utf-8-sig") as fh:
        return {r["match_id"]: r for r in csv.DictReader(fh) if r.get("match_id")}


def merge_and_save(path: Path, fresh: list[dict[str, Any]]) -> list[dict[str, Any]]:
    old = read_existing_csv(path)
    for row in fresh:
        prior = old.get(row["match_id"], {})
        # User-editable fields survive future scrapes.
        row["notes"] = prior.get("notes", "")
        row["tags"] = prior.get("tags", "")
        old[row["match_id"]] = {k: row.get(k, "") for k in RAW_FIELDS}
    rows = list(old.values())
    temp = path.with_suffix(".csv.tmp")
    with temp.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=RAW_FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    temp.replace(path)
    return rows


def triples(values: list[float]) -> tuple[float | str, float | str, float | str]:
    if not values:
        return "", "", ""
    return round(statistics.mean(values), 2), min(values), max(values)


def team_stats(rows: list[dict[str, Any]]) -> list[list[Any]]:
    buckets: dict[tuple[str, str], dict[str, list[float]]] = defaultdict(lambda: defaultdict(list))
    for r in rows:
        league = str(r.get("league", ""))
        for side, opponent in (("home", "away"), ("away", "home")):
            team = str(r.get(f"{side}_team", ""))
            if not team:
                continue
            b = buckets[(league, team)]
            for q in range(1, 5):
                v = number(r.get(f"{side}_q{q}"))
                if v is not None:
                    b[f"q{q}"].append(v)
            total = number(r.get(f"{side}_total"))
            opp = number(r.get(f"{opponent}_total"))
            if total is not None:
                b["total"].append(total)
            if opp is not None:
                b["allowed"].append(opp)

    header = ["League", "Team", "Games"]
    for label in ["Q1", "Q2", "Q3", "Q4", "Total", "Allowed"]:
        header += [f"{label} Avg", f"{label} Min", f"{label} Max"]
    output = [header]
    for (league, team), b in sorted(buckets.items()):
        line: list[Any] = [league, team, len(b["total"])]
        for key in ["q1", "q2", "q3", "q4", "total", "allowed"]:
            line.extend(triples(b[key]))
        output.append(line)
    return output


def combined_stats(rows: list[dict[str, Any]]) -> list[list[Any]]:
    buckets: dict[str, dict[str, list[float]]] = defaultdict(lambda: defaultdict(list))
    for r in rows:
        league = str(r.get("league", ""))
        b = buckets[league]
        for q in range(1, 5):
            h, a = number(r.get(f"home_q{q}")), number(r.get(f"away_q{q}"))
            if h is not None and a is not None:
                b[f"q{q}"].append(h + a)
        ht, at = number(r.get("home_total")), number(r.get("away_total"))
        if ht is not None and at is not None:
            b["game"].append(ht + at)
    header = ["League", "Games"]
    for label in ["Combined Q1", "Combined Q2", "Combined Q3", "Combined Q4", "Combined Game"]:
        header += [f"{label} Avg", f"{label} Min", f"{label} Max"]
    output = [header]
    for league, b in sorted(buckets.items()):
        line: list[Any] = [league, len(b["game"])]
        for key in ["q1", "q2", "q3", "q4", "game"]:
            line.extend(triples(b[key]))
        output.append(line)
    return output


def write_sheet(ws, rows: list[list[Any]]) -> None:
    for row in rows:
        ws.append(row)
    if ws.max_row:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        width = min(35, max(10, max(len(str(ws.cell(r, col).value or "")) for r in range(1, ws.max_row + 1)) + 2))
        ws.column_dimensions[get_column_letter(col)].width = width


def export_xlsx(path: Path, rows: list[dict[str, Any]], config: dict[str, Any]) -> None:
    wb = Workbook()
    games = wb.active
    games.title = "Games"
    write_sheet(games, [RAW_FIELDS] + [[r.get(k, "") for k in RAW_FIELDS] for r in rows])
    write_sheet(wb.create_sheet("Team Quarter Stats"), team_stats(rows))
    write_sheet(wb.create_sheet("League Combined Stats"), combined_stats(rows))
    settings = wb.create_sheet("Settings")
    settings_rows = [["Setting", "Value"], ["Generated UTC", datetime.now(timezone.utc).isoformat(timespec="seconds")]]
    settings_rows += [["League URL", u] for u in config.get("league_urls", [])]
    settings_rows += [["Note", "Edit notes/tags in games.csv; those columns are preserved on future scrapes."]]
    write_sheet(settings, settings_rows)
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Scrape Flashscore basketball results and calculate quarter stats.")
    parser.add_argument("--config", default=str(ROOT / "config.json"))
    parser.add_argument("--rebuild-only", action="store_true", help="Rebuild the workbook from existing CSV without scraping.")
    args = parser.parse_args()
    config_path = Path(args.config).resolve()
    config = json.loads(config_path.read_text(encoding="utf-8"))
    out = (config_path.parent / config.get("output_directory", "data")).resolve()
    out.mkdir(parents=True, exist_ok=True)
    csv_path, xlsx_path = out / "games.csv", out / "basketball_data.xlsx"
    fresh = [] if args.rebuild_only else scrape(config)
    rows = merge_and_save(csv_path, fresh)
    for filename, table in (("team_stats.csv", team_stats(rows)), ("league_stats.csv", combined_stats(rows))):
        with (out / filename).open("w", newline="", encoding="utf-8-sig") as fh:
            csv.writer(fh).writerows(table)
    export_xlsx(xlsx_path, rows, config)
    print(f"Saved {len(rows)} unique games to {csv_path}")
    print(f"Workbook: {xlsx_path}")


if __name__ == "__main__":
    main()
